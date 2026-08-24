import pandas as pd
from openpyxl.styles import PatternFill


# ======================================================
# LOAD HEREDITARY GENES
# ======================================================
def load_genes(file_path):

    df = pd.read_excel(file_path, dtype=str)
    df.columns = df.columns.str.strip()

    col_name = "Hereditary (germline) - 158 genes"

    if col_name not in df.columns:
        raise ValueError(f"Column '{col_name}' not found in gene file")

    return set(
        df[col_name]
        .dropna()
        .astype(str)
        .str.strip()
        .str.upper()
    )


# ======================================================
# SORT DATAFRAME (MATCHED ROWS FIRST)
# ======================================================
def sort_df(df, hereditary_genes):

    cols = df.columns.tolist()
    symbol_cols = [c for c in cols if c in ["SYMBOL", "SYMBOL_X"]]

    if not symbol_cols:
        return df

    def check_row(row):
        for col in symbol_cols:
            gene = str(row[col]).strip().upper()
            if gene in hereditary_genes:
                return True
        return False

    df["_MATCH"] = df.apply(check_row, axis=1)

    df = df.sort_values(by="_MATCH", ascending=False)

    return df.drop(columns="_MATCH")


# ======================================================
# HIGHLIGHT ONE SHEET
# ======================================================
def highlight_sheet(writer, sheet_name, hereditary_genes, fill_color):

    ws = writer.book[sheet_name]

    header = [cell.value for cell in ws[1]]

    symbol_cols = []
    if "SYMBOL" in header:
        symbol_cols.append(header.index("SYMBOL") + 1)
    if "SYMBOL_X" in header:
        symbol_cols.append(header.index("SYMBOL_X") + 1)

    if not symbol_cols:
        return

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):

        match = False

        for col_idx in symbol_cols:
            val = ws.cell(row=row[0].row, column=col_idx).value
            gene = str(val).strip().upper() if val else ""

            if gene in hereditary_genes:
                match = True
                break

        if match:
            for cell in row:
                cell.fill = fill_color


# ======================================================
# APPLY (SORT + HIGHLIGHT)
# ======================================================
def apply_highlighting(writer, hereditary_genes):

    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    target_sheets = [
        "CS_Pathogenic",
        "uncertain_significance",
        "Others",
        "ClinGen_Matched",
        "ClinGen_Unmatched"
    ]

    # --------------------------
    # STEP 1: SORT BEFORE WRITE
    # --------------------------
    for sheet in target_sheets:

        if sheet not in writer.book.sheetnames:
            continue

        ws = writer.book[sheet]

        data = ws.values
        columns = next(data, None)

        # ✅ Fix: handle empty sheets
        if columns is None:
            continue

        df = pd.DataFrame(data, columns=columns)

        df_sorted = sort_df(df, hereditary_genes)

        # Rewrite sheet
        ws.delete_rows(1, ws.max_row)

        ws.append(list(df_sorted.columns))

        for row in df_sorted.itertuples(index=False):
            ws.append(list(row))

    # --------------------------
    # STEP 2: APPLY HIGHLIGHT
    # --------------------------
    for sheet in target_sheets:

        if sheet not in writer.book.sheetnames:
            continue

        highlight_sheet(writer, sheet, hereditary_genes, red)
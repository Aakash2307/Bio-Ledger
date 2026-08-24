import pandas as pd
from openpyxl.styles import PatternFill

from config import (
    SHEET_NAME 
)

# ======================================================
# LOAD HEREDITARY GENES
# ======================================================
def load_genes(file_path):

    df = pd.read_excel(file_path, dtype=str)
    df.columns = df.columns.str.strip()

    return set(
        df["Hereditary (germline) - 158 genes"]
        .dropna().astype(str).str.strip().str.upper()
    )


# ======================================================
# LOAD DISEASE GENE SETS
# ======================================================
def load_disease_gene_sets(file_path):

    df = pd.read_excel(file_path, dtype=str)
    df.columns = df.columns.str.strip()

    disease_cols = [
        "Cardiac - 565 genes",
        "Diabetes_pharmgkb_mito genes - 318 genes",
        "NDD - 890 genes",
        "Endometriosis - 53 genes",
        "Neurodevelopmental - 1733 genes",
    ]

    return {
        col: set(df[col].dropna().astype(str).str.strip().str.upper())
        for col in disease_cols if col in df.columns
    }


# ======================================================
# SORT BY COLOR PRIORITY
# ======================================================
def sort_df(df, hereditary_genes, disease_gene_sets):

    symbol_cols = [c for c in df.columns if c in ["SYMBOL", "SYMBOL_X"]]
    if not symbol_cols:
        return df

    def get_priority(row):

        matched = set()
        hereditary = False

        for col in symbol_cols:
            gene = str(row[col]).strip().upper()
            if not gene:
                continue

            if gene in hereditary_genes:
                hereditary = True

            for d, gset in disease_gene_sets.items():
                if gene in gset:
                    matched.add(d)

        if hereditary:
            return 6
        if len(matched) == 1:
            order = {
                "Cardiac - 565 genes": 5,
                "Neurodevelopmental - 1733 genes": 4,
                "Endometriosis - 53 genes": 3,
                "Diabetes_pharmgkb_mito genes - 318 genes": 2,
                "NDD - 890 genes": 2
            }
            return order.get(list(matched)[0], 1)
        if len(matched) > 1:
            return 1
        return 0

    df["_P"] = df.apply(get_priority, axis=1)
    df = df.sort_values("_P", ascending=False)
    return df.drop(columns="_P")


# ======================================================
# APPLY ALL HIGHLIGHTING
# ======================================================
def apply_highlighting(writer, hereditary_genes, disease_gene_sets):

    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    fills = {
        "Cardiac - 565 genes": PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid"),
        "Diabetes_pharmgkb_mito genes - 318 genes": PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid"),
        "NDD - 890 genes": PatternFill(start_color="C0C0C0", end_color="C0C0C0", fill_type="solid"),
        "Endometriosis - 53 genes": PatternFill(start_color="FF69B4", end_color="FF69B4", fill_type="solid"),
        "Neurodevelopmental - 1733 genes": PatternFill(start_color="87CEEB", end_color="87CEEB", fill_type="solid"),
    }

    multi_fill = PatternFill(start_color="E6CCFF", end_color="E6CCFF", fill_type="solid")

    sheets = [
        SHEET_NAME,
        "CS_Pathogenic",
        "uncertain_significance",
        "Others",
        "ClinGen_Matched",
        "ClinGen_Unmatched"
    ]

    for sheet in sheets:
        if sheet not in writer.book.sheetnames:
            continue

        ws = writer.book[sheet]

        data = ws.values
        cols = next(data, None)
        if cols is None:
            continue

        df = pd.DataFrame(data, columns=cols)
        df = sort_df(df, hereditary_genes, disease_gene_sets)

        ws.delete_rows(1, ws.max_row)
        ws.append(list(df.columns))

        for r in df.itertuples(index=False):
            ws.append(list(r))

        header = [c.value for c in ws[1]]

        symbol_cols = []
        if "SYMBOL" in header:
            symbol_cols.append(header.index("SYMBOL") + 1)
        if "SYMBOL_X" in header:
            symbol_cols.append(header.index("SYMBOL_X") + 1)

        match_col = len(header) + 1
        ws.cell(row=1, column=match_col, value="MATCH_SOURCE")

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):

            matched = set()
            hereditary = False

            for col_idx in symbol_cols:
                gene = str(ws.cell(row=row[0].row, column=col_idx).value).strip().upper()

                if gene in hereditary_genes:
                    hereditary = True

                for d, gset in disease_gene_sets.items():
                    if gene in gset:
                        matched.add(d)

            if matched:
                ws.cell(row=row[0].row, column=match_col).value = ", ".join(sorted(matched))

            if len(matched) == 1:
                fill = fills[list(matched)[0]]
            elif len(matched) > 1:
                fill = multi_fill
            elif hereditary:
                fill = red
            else:
                continue

            for cell in row:
                cell.fill = fill

    # =======================
    # LEGEND SHEET
    # =======================
    legend = writer.book.create_sheet("Color codes ")

    legend_data = [
        ("Red", "Hereditary genes"),
        ("Orange", "Cardiac"),
        ("Sky Blue", "Neurodevelopmental"),
        ("Pink", "Endometriosis"),
        ("Green", "Diabetes"),
        ("Grey", "NDD"),
        ("Light Purple", "Multi-match")
    ]

    for i, (color, desc) in enumerate(legend_data, start=1):
        legend.cell(row=i, column=1, value=color)
        legend.cell(row=i, column=2, value=desc)

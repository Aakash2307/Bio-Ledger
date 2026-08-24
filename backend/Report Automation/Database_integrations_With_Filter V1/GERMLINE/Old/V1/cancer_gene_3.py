import pandas as pd
from openpyxl.styles import PatternFill

from config import (
    SHEET_NAME
)

# ======================================================
# LOAD ALL GENE SETS
# ======================================================
def load_disease_gene_sets(file_path):

    df = pd.read_excel(file_path, dtype=str)
    df.columns = df.columns.str.strip()

    disease_cols = [
        "Hereditary (germline) - 158 genes",
        "Cardiac - 565 genes",
        "Diabetes_pharmgkb_mito genes - 318 genes",
        "NDD - 890 genes",
        "Endometriosis - 53 genes",
        "Neurodevelopmental - 1733 genes",
    ]

    return {
        col: set(df[col].dropna().astype(str).str.strip().str.upper())
        for col in disease_cols if col in df.columns
    }


# ======================================================
# SORT BY PRIORITY
# ======================================================
def sort_df(df, disease_gene_sets):

    symbol_cols = [c for c in df.columns if c in ["SYMBOL", "SYMBOL_X"]]

    if not symbol_cols:
        return df

    priority_order = {
        "Hereditary (germline) - 158 genes": 6,
        "Cardiac - 565 genes": 5,
        "Neurodevelopmental - 1733 genes": 4,
        "Endometriosis - 53 genes": 3,
        "Diabetes_pharmgkb_mito genes - 318 genes": 2,
        "NDD - 890 genes": 1,
    }

    def get_priority(row):

        matched = set()

        for col in symbol_cols:

            gene = str(row[col]).strip().upper()

            if not gene:
                continue

            for disease, gset in disease_gene_sets.items():

                if gene in gset:
                    matched.add(disease)

        # MULTI MATCH
        if len(matched) > 1:
            return 7

        # SINGLE MATCH
        elif len(matched) == 1:
            disease = list(matched)[0]
            return priority_order.get(disease, 0)

        return 0

    df["_P"] = df.apply(get_priority, axis=1)

    df = df.sort_values("_P", ascending=False)

    return df.drop(columns="_P")


# ======================================================
# APPLY HIGHLIGHTING
# ======================================================
def apply_highlighting(writer, disease_gene_sets):

    fills = {
        "Hereditary (germline) - 158 genes":
            PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),

        "Cardiac - 565 genes":
            PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid"),

        "Diabetes_pharmgkb_mito genes - 318 genes":
            PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid"),

        "NDD - 890 genes":
            PatternFill(start_color="C0C0C0", end_color="C0C0C0", fill_type="solid"),

        "Endometriosis - 53 genes":
            PatternFill(start_color="FF69B4", end_color="FF69B4", fill_type="solid"),

        "Neurodevelopmental - 1733 genes":
            PatternFill(start_color="87CEEB", end_color="87CEEB", fill_type="solid"),
    }

    # MULTI MATCH COLOR
    multi_fill = PatternFill(
        start_color="E6CCFF",
        end_color="E6CCFF",
        fill_type="solid"
    )

    sheets = [
        SHEET_NAME,
        "CS_Pathogenic",
        "uncertain_significance",
        "Others",
        "ClinGen_Matched",
        "ClinGen_Unmatched"
    ]

    for sheet in sheets:

        if sheet not in writer.book.sheetnames:
            continue

        ws = writer.book[sheet]

        data = ws.values

        cols = next(data, None)

        if cols is None:
            continue

        df = pd.DataFrame(data, columns=cols)

        # SORT DATAFRAME
        df = sort_df(df, disease_gene_sets)

        # CLEAR SHEET
        ws.delete_rows(1, ws.max_row)

        # WRITE HEADER
        ws.append(list(df.columns))

        # WRITE DATA
        for r in df.itertuples(index=False):
            ws.append(list(r))

        header = [c.value for c in ws[1]]

        symbol_cols = []

        if "SYMBOL" in header:
            symbol_cols.append(header.index("SYMBOL") + 1)

        if "SYMBOL_X" in header:
            symbol_cols.append(header.index("SYMBOL_X") + 1)

        # ADD MATCH SOURCE COLUMN
        match_col = len(header) + 1

        ws.cell(row=1, column=match_col, value="MATCH_SOURCE")

        # PROCESS ROWS
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):

            matched = set()

            for col_idx in symbol_cols:

                gene = str(
                    ws.cell(row=row[0].row, column=col_idx).value
                ).strip().upper()

                for disease, gset in disease_gene_sets.items():

                    if gene in gset:
                        matched.add(disease)

            # WRITE MATCH SOURCE
            if matched:

                ws.cell(
                    row=row[0].row,
                    column=match_col
                ).value = ", ".join(sorted(matched))

            # SELECT COLOR
            if len(matched) > 1:

                fill = multi_fill

            elif len(matched) == 1:

                disease = list(matched)[0]

                fill = fills[disease]

            else:
                continue

            # APPLY COLOR TO ENTIRE ROW
            for cell in row:
                cell.fill = fill

    # ==================================================
    # LEGEND SHEET
    # ==================================================
    legend = writer.book.create_sheet("Color codes")

    legend_data = [
        ("Red", "Hereditary"),
        ("Orange", "Cardiac"),
        ("Sky Blue", "Neurodevelopmental"),
        ("Pink", "Endometriosis"),
        ("Green", "Diabetes"),
        ("Grey", "NDD"),
        ("Light Purple", "Multi-match")
    ]

    for i, (color, desc) in enumerate(legend_data, start=1):

        legend.cell(row=i, column=1, value=color)

        legend.cell(row=i, column=2, value=desc)
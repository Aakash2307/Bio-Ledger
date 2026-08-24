# ======================================================
# disease_sheet_generator.py
# ======================================================

import pandas as pd
from openpyxl.styles import PatternFill

# ======================================================
# LOAD DISEASE GENE SETS
# ======================================================
def load_disease_gene_sets(file_path):

    df = pd.read_excel(file_path, dtype=str)

    df.columns = df.columns.str.strip()

    disease_cols = [
        "Hereditary (germline) - 158 genes",
        "Cardiac - 565 genes",
        "Diabetes_pharmgkb_mito genes - 318 genes",
        "NDD - 890 genes",
        "Endometriosis - 53 genes",
        "Neurodevelopmental - 1733 genes",
    ]

    disease_gene_sets = {}

    for col in disease_cols:

        if col not in df.columns:
            continue

        disease_gene_sets[col] = set(
            df[col]
            .dropna()
            .astype(str)
            .str.strip()
            .str.upper()
        )

    return disease_gene_sets

# ======================================================
# CREATE DISEASE SHEETS
# ======================================================
def create_disease_sheets(
    writer,
    disease_gene_sets
):

    # ==================================================
    # CLINICAL SIGNIFICANCE COLORS
    # ==================================================
    source_fills = {

        "CS_Pathogenic":
            PatternFill(
                start_color="FFC7CE",
                end_color="FFC7CE",
                fill_type="solid"
            ),

        "uncertain_significance":
            PatternFill(
                start_color="FFF3CD",
                end_color="FFF3CD",
                fill_type="solid"
            ),

        "Others":
            PatternFill(
                start_color="E2E3E5",
                end_color="E2E3E5",
                fill_type="solid"
            ),
    }

    # ==================================================
    # SOURCE SHEETS
    # ==================================================
    source_sheets = [
        "CS_Pathogenic",
        "uncertain_significance",
        "Others"
    ]

    # ==================================================
    # PROCESS EACH DISEASE
    # ==================================================
    for disease_name, gene_set in disease_gene_sets.items():

        disease_rows = []

        # ==============================================
        # COLLECT ROWS FROM SOURCE SHEETS
        # ==============================================
        for source_sheet in source_sheets:

            if source_sheet not in writer.book.sheetnames:
                continue

            ws = writer.book[source_sheet]

            data = ws.values

            cols = next(data, None)

            if cols is None:
                continue

            df = pd.DataFrame(
                data,
                columns=cols
            )

            # FIND SYMBOL COLUMNS
            symbol_cols = [
                c for c in df.columns
                if c in ["SYMBOL", "SYMBOL_X"]
            ]

            if not symbol_cols:
                continue

            # ==========================================
            # MATCH GENES
            # ==========================================
            matched_mask = pd.Series(
                False,
                index=df.index
            )

            for col in symbol_cols:

                matched_mask |= (
                    df[col]
                    .fillna("")
                    .astype(str)
                    .str.strip()
                    .str.upper()
                    .isin(gene_set)
                )

            matched_df = df[matched_mask].copy()

            if matched_df.empty:
                continue

            matched_df["SOURCE_SHEET"] = source_sheet

            matched_df["MATCHED_DISEASE"] = disease_name

            disease_rows.append(matched_df)

        # ==============================================
        # SKIP EMPTY DISEASE SHEETS
        # ==============================================
        if not disease_rows:
            continue

        # ==============================================
        # COMBINE ROWS
        # ==============================================
        final_df = pd.concat(
            disease_rows,
            ignore_index=True
        )

        # ==============================================
        # CREATE CLEAN SHEET NAME
        # ==============================================
        clean_sheet_name = (
            disease_name
            .replace(" - ", "_")
            .replace(" ", "_")
        )[:31]

        # ==============================================
        # WRITE SHEET
        # ==============================================
        final_df.to_excel(
            writer,
            sheet_name=clean_sheet_name,
            index=False
        )

        # ==============================================
        # APPLY COLORS
        # ==============================================
        ws = writer.book[clean_sheet_name]

        header = [
            cell.value
            for cell in ws[1]
        ]

        try:
            source_col_idx = (
                header.index("SOURCE_SHEET") + 1
            )

        except ValueError:
            continue

        for row in ws.iter_rows(
            min_row=2,
            max_row=ws.max_row
        ):

            source_value = ws.cell(
                row=row[0].row,
                column=source_col_idx
            ).value

            if source_value not in source_fills:
                continue

            fill = source_fills[source_value]

            for cell in row:
                cell.fill = fill

# ======================================================
# disease_sheet_generator.py
# ======================================================

import logging
import pandas as pd
from collections import defaultdict
from openpyxl.styles import PatternFill

# ======================================================
# LOGGING CONFIG
# ======================================================
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s"
)

# ======================================================
# LOAD DISEASE GENE SETS
# ======================================================
def load_disease_gene_sets(file_path):

    logging.info("Loading disease gene sets...")

    df = pd.read_excel(file_path, dtype=str)

    df.columns = df.columns.str.strip()

    # ==================================================
    # AUTO DETECT DISEASE COLUMNS
    # ==================================================
    disease_cols = [
        col for col in df.columns
        if "gene" in col.lower()
    ]

    logging.info(f"Detected disease columns: {disease_cols}")

    disease_gene_sets = {}

    # ==================================================
    # BUILD DISEASE → GENE SET
    # ==================================================
    for col in disease_cols:

        genes = set(
            df[col]
            .dropna()
            .astype(str)
            .str.strip()
            .str.upper()
        )

        disease_gene_sets[col] = genes

        logging.info(
            f"{col} -> {len(genes)} genes loaded"
        )

    return disease_gene_sets


# ======================================================
# BUILD REVERSE GENE MAPPING
# GENE -> ALL DISEASES
# ======================================================
def build_gene_to_disease_map(disease_gene_sets):

    gene_to_diseases = defaultdict(set)

    for disease_name, gene_set in disease_gene_sets.items():

        for gene in gene_set:

            gene_to_diseases[gene].add(disease_name)

    logging.info(
        f"Reverse mapping built for "
        f"{len(gene_to_diseases)} genes"
    )

    return gene_to_diseases


# ======================================================
# CREATE DISEASE SHEETS
# ======================================================
def create_disease_sheets(
    writer,
    disease_gene_sets
):

    # ==================================================
    # BUILD REVERSE MAPPING
    # ==================================================
    gene_to_diseases = build_gene_to_disease_map(
        disease_gene_sets
    )

    # ==================================================
    # CLINICAL SIGNIFICANCE COLORS
    # ==================================================
    source_fills = {

        "CS_Pathogenic":
            PatternFill(
                start_color="FFC7CE",
                end_color="FFC7CE",
                fill_type="solid"
            ),

        "uncertain_significance":
            PatternFill(
                start_color="FFF3CD",
                end_color="FFF3CD",
                fill_type="solid"
            ),

        "Others":
            PatternFill(
                start_color="E2E3E5",
                end_color="E2E3E5",
                fill_type="solid"
            ),
    }

    # ==================================================
    # SOURCE SHEETS
    # ==================================================
    source_sheets = [
        sheet for sheet in writer.book.sheetnames
        if sheet in [
            "CS_Pathogenic",
            "uncertain_significance",
            "Others"
        ]
    ]

    logging.info(
        f"Source sheets found: {source_sheets}"
    )

    # ==================================================
    # PROCESS EACH DISEASE
    # ==================================================
    for disease_name, gene_set in disease_gene_sets.items():

        logging.info(
            f"Processing disease: {disease_name}"
        )

        disease_rows = []

        # ==============================================
        # PROCESS SOURCE SHEETS
        # ==============================================
        for source_sheet in source_sheets:

            ws = writer.book[source_sheet]

            data = ws.values

            cols = next(data, None)

            if cols is None:
                continue

            df = pd.DataFrame(
                data,
                columns=cols
            )

            # ==========================================
            # FIND SYMBOL COLUMNS
            # ==========================================
            symbol_cols = [
                c for c in df.columns
                if str(c).strip().upper()
                in ["SYMBOL", "SYMBOL_X"]
            ]

            if not symbol_cols:

                logging.warning(
                    f"No SYMBOL column found in "
                    f"{source_sheet}"
                )

                continue

            # ==========================================
            # MATCH GENES
            # ==========================================
            matched_mask = pd.Series(
                False,
                index=df.index
            )

            for col in symbol_cols:

                matched_mask |= (
                    df[col]
                    .fillna("")
                    .astype(str)
                    .str.strip()
                    .str.upper()
                    .isin(gene_set)
                )

            matched_df = df[matched_mask].copy()

            if matched_df.empty:
                continue

            # ==========================================
            # DETERMINE ALL MATCHED DISEASES
            # ==========================================
            def get_all_diseases(row):

                all_diseases = set()

                for col in symbol_cols:

                    gene = (
                        str(row[col])
                        .strip()
                        .upper()
                    )

                    if gene in gene_to_diseases:

                        all_diseases.update(
                            gene_to_diseases[gene]
                        )

                return ", ".join(
                    sorted(all_diseases)
                )

            matched_df[
                "ALL_MATCHED_DISEASES"
            ] = matched_df.apply(
                get_all_diseases,
                axis=1
            )

            # ==========================================
            # ADD SOURCE SHEET
            # ==========================================
            matched_df[
                "SOURCE_SHEET"
            ] = source_sheet

            # ==========================================
            # ADD PRIMARY DISEASE
            # ==========================================
            matched_df[
                "PRIMARY_MATCHED_DISEASE"
            ] = disease_name

            logging.info(
                f"{source_sheet} -> "
                f"{len(matched_df)} rows matched"
            )

            disease_rows.append(matched_df)

        # ==============================================
        # SKIP EMPTY SHEETS
        # ==============================================
        if not disease_rows:

            logging.warning(
                f"No rows matched for "
                f"{disease_name}"
            )

            continue

        # ==============================================
        # COMBINE RESULTS
        # ==============================================
        final_df = pd.concat(
            disease_rows,
            ignore_index=True
        )

        # ==============================================
        # REMOVE DUPLICATES
        # ==============================================
        final_df.drop_duplicates(
            inplace=True
        )



        
        # ==============================================
        # MOVE METADATA COLUMNS TO END
        # ==============================================
        metadata_cols = [
            "SOURCE_SHEET",
            "PRIMARY_MATCHED_DISEASE",
            "ALL_MATCHED_DISEASES"
        ]

        existing_metadata_cols = [
            col for col in metadata_cols
            if col in final_df.columns
        ]

        other_cols = [
            col for col in final_df.columns
            if col not in existing_metadata_cols
        ]

        final_df = final_df[
            other_cols + existing_metadata_cols
        ]



        # ==============================================
        # CLEAN SHEET NAME
        # ==============================================
        clean_sheet_name = (
            disease_name
            .replace(" - ", "_")
            .replace(" ", "_")
        )[:31]

        # ==============================================
        # WRITE SHEET
        # ==============================================
        final_df.to_excel(
            writer,
            sheet_name=clean_sheet_name,
            index=False
        )

        logging.info(
            f"Written sheet: {clean_sheet_name}"
        )

        # ==============================================
        # APPLY COLORS
        # ==============================================
        ws = writer.book[clean_sheet_name]

        header = [
            cell.value
            for cell in ws[1]
        ]

        try:
            source_col_idx = (
                header.index("SOURCE_SHEET") + 1
            )

        except ValueError:
            continue

        for row in ws.iter_rows(
            min_row=2,
            max_row=ws.max_row
        ):

            source_value = ws.cell(
                row=row[0].row,
                column=source_col_idx
            ).value

            if source_value not in source_fills:
                continue

            fill = source_fills[source_value]

            for cell in row:
                cell.fill = fill

    logging.info(
        "Disease sheet generation completed."
    )


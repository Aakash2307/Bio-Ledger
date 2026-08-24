import os
import sys

# 🔥 MUST be first
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

import pandas as pd  # type: ignore
from tqdm import tqdm # type: ignore
from openpyxl import load_workbook # type: ignore
from openpyxl.utils.dataframe import dataframe_to_rows # type: ignore

from add_filter import run_all_filters
from somatic_highlighter2 import (
    load_somatic_genes,
    load_disease_gene_sets,   # ✅ FIXED
    highlight_excel_file,
    apply_disease_highlighting
)

# ======================================================
# CONFIG
# ======================================================
from config import ( # type: ignore
    SOMATIC_OUTPUT,
    SOMAMUT_REF,
    GENE_LIST_FILE,
    SHEET_NAME,
    SOMATIC_INPUT
)

# ============================================
# Helper: Extract sample prefix (PD003 → PD0003)
# Mirrors the logic in pre_filter.py so matching stays consistent
# ============================================
def get_sample_prefix(filename):
    prefix = filename.split("_")[0]

    if prefix.startswith("PD"):
        num = prefix[2:]
        if num.isdigit():
            return f"PD{int(num):04d}"

    return prefix


# ============================================
# Restore the ORIGINAL (pre germline-filter) Variations sheet
# into the final _Integrated_Somamut.xlsx output.
#
# Deletes whatever Variations sheet currently exists in output_file
# (which has already been through pre_filter.py's germline-location
# filtering AND add_filter.py's CADD/TAF/AF filters) and replaces it
# with the raw Variations sheet read straight from SOMATIC_INPUT —
# i.e. the file as it looked before pre_filter.py ever touched it.
# ============================================
def restore_original_variations(output_file, sample_prefix):

    # sample_prefix here is the raw output-folder name from Master_Script_7.py
    # (e.g. "PD003", unpadded). Normalize it the same way pre_filter.py
    # normalizes filenames, so it correctly matches files in SOMATIC_INPUT
    # (e.g. "PD0003_somatic.xlsx").
    normalized_prefix = get_sample_prefix(sample_prefix)

    raw_files = [
        f for f in os.listdir(SOMATIC_INPUT)
        if f.lower().endswith((".xlsx", ".xls"))
        and get_sample_prefix(f) == normalized_prefix
    ]

    if not raw_files:
        print(f"⚠️ No raw somatic file found in SOMATIC_INPUT for sample {sample_prefix} (normalized: {normalized_prefix}) — Variations sheet NOT restored")
        return

    raw_path = os.path.join(SOMATIC_INPUT, raw_files[0])

    try:
        original_variations_df = pd.read_excel(
            raw_path,
            sheet_name=SHEET_NAME,
            dtype=str
        ).fillna("")
    except Exception as e:
        print(f"❌ Failed to read original Variations sheet from {raw_path}: {e}")
        return

    wb = load_workbook(output_file)

    if SHEET_NAME in wb.sheetnames:
        original_index = wb.sheetnames.index(SHEET_NAME)
        del wb[SHEET_NAME]
    else:
        original_index = 0

    ws = wb.create_sheet(SHEET_NAME, original_index)

    for row in dataframe_to_rows(original_variations_df, index=False, header=True):
        ws.append(row)

    wb.save(output_file)

    print(f"✔ Restored original (pre-filter) Variations sheet from: {raw_path}")


# ============================================
# Helper: Ensure unique column names
# ============================================
def make_unique_columns(columns):
    seen = {}
    new_cols = []
    for col in columns:
        col_clean = str(col).strip().upper()
        if col_clean not in seen:
            seen[col_clean] = 1
            new_cols.append(col_clean)
        else:
            seen[col_clean] += 1
            new_cols.append(f"{col_clean}_{seen[col_clean]}")
    return new_cols


# ============================================
# Helper: Resolve POS / REF / ALT columns
# ============================================
def resolve_variant_columns(df):
    cols = set(df.columns)
    if {"POS", "REF", "ALT"}.issubset(cols):
        return "POS", "REF", "ALT"
    if {"POS_X", "REF_X", "ALT_X"}.issubset(cols):
        return "POS_X", "REF_X", "ALT_X"
    return None, None, None


print("\n==============================")
print("   SOMAMUT FILTER STARTED     ")
print("==============================\n")

# ============================================
# LOAD GENE DATA
# ============================================
somatic_genes = load_somatic_genes(GENE_LIST_FILE)
disease_gene_sets = load_disease_gene_sets(GENE_LIST_FILE)  # ✅ FIXED

print(f"✔ Somatic genes loaded: {len(somatic_genes)}")
print(f"✔ Disease gene sets loaded: {len(disease_gene_sets)}")


# ============================================
# LOOP OVER SAMPLE OUTPUT FOLDERS
# ============================================
for sample_folder in os.listdir(SOMATIC_OUTPUT):

    sample_dir = os.path.join(SOMATIC_OUTPUT, sample_folder)

    if not os.path.isdir(sample_dir):
        continue

    integrated_files = [
        f for f in os.listdir(sample_dir)
        if f.endswith("_Integrated.xlsx")
    ]

    if not integrated_files:
        continue

    sample_file = os.path.join(sample_dir, integrated_files[0])

    output_file = os.path.join(
        sample_dir,
        integrated_files[0].replace(
            "_Integrated.xlsx",
            "_Integrated_Somamut.xlsx"
        )
    )

    print(f"\n✔ Sample file detected: {sample_file}")

    xls = pd.ExcelFile(sample_file)
    all_sheets = xls.sheet_names
    print(f"✔ Sheets found: {all_sheets}")

    SKIP_VARIATIONS = SHEET_NAME
    SUMMARY_SHEET = "RUN_SUMMARY"

    # ============================================
    # LOAD REFERENCE CSVs
    # ============================================
    reference_set = set()

    ref_files = [
        f for f in os.listdir(SOMAMUT_REF)
        if f.lower().endswith(".csv")
    ]

    print(f"✔ Reference CSV files found: {len(ref_files)}")
    print("⏳ Indexing reference variants...")

    for file in tqdm(ref_files, desc="Processing reference CSVs"):

        ref_path = os.path.join(SOMAMUT_REF, file)
        ref_df = pd.read_csv(ref_path, dtype=str).fillna("")
        ref_df.columns = make_unique_columns(ref_df.columns)

        if not {"POS", "REF", "ALT"}.issubset(ref_df.columns):
            continue

        ref_df["POS"] = ref_df["POS"].str.strip()
        ref_df["REF"] = ref_df["REF"].str.strip().str.upper()
        ref_df["ALT"] = ref_df["ALT"].str.strip().str.upper()

        reference_set.update(
            zip(ref_df["POS"], ref_df["REF"], ref_df["ALT"])
        )

    print(f"✔ Total reference variants indexed: {len(reference_set)}")

    # after_counts is keyed by the NORMALIZED sheet name (sheet.strip().upper())
    # so lookups during the Run_Summary update below are case-insensitive and
    # don't depend on row order. See requirement: dict instead of list.
    after_counts = {}
    run_summary_df = None

    # ============================================
    # WRITE OUTPUT
    # ============================================
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:

        for sheet in all_sheets:

            sheet_norm = sheet.strip().upper()
            print(f"\n🔍 Processing sheet: {sheet}")

            df = pd.read_excel(sample_file, sheet_name=sheet, dtype=str).fillna("")
            df.columns = make_unique_columns(df.columns)

            # Preserve summary
            if sheet_norm == SUMMARY_SHEET:
                run_summary_df = df
                continue

            pos_col, ref_col, alt_col = resolve_variant_columns(df)

            # INCLUDE VARIATIONS (no skipping)
            if not pos_col or sheet_norm == SKIP_VARIATIONS:
                df.to_excel(writer, sheet_name=sheet, index=False)
                after_counts[sheet_norm] = len(df)
                continue

            df[pos_col] = df[pos_col].str.strip()
            df[ref_col] = df[ref_col].str.strip().str.upper()
            df[alt_col] = df[alt_col].str.strip().str.upper()

            df["_KEY"] = list(zip(df[pos_col], df[ref_col], df[alt_col]))

            filtered_df = df.loc[
                ~df["_KEY"].isin(reference_set)
            ].drop(columns="_KEY")

            print(f"✔ {sheet} | Before: {len(df)} | After: {len(filtered_df)}")

            after_counts[sheet_norm] = len(filtered_df)
            filtered_df.to_excel(writer, sheet_name=sheet, index=False)

        # ============================================
        # UPDATE RUN SUMMARY
        # Matches each Run_Summary row to after_counts by NORMALIZED
        # database name (strip + upper) instead of positional row_idx,
        # so the mapping is correct regardless of sheet order.
        # ============================================
        if run_summary_df is not None:

            print("\n📝 Updating Run_Summary")

            # Ensure 4 columns exist
            while run_summary_df.shape[1] < 4:
                run_summary_df[f"COL_{run_summary_df.shape[1] + 1}"] = ""

            cols = list(run_summary_df.columns)
            cols[0] = "Database"
            cols[1] = "Before"
            cols[2] = "After_Filter"
            cols[3] = "Excluded"
            run_summary_df.columns = cols

            # Remove TOTAL row if already present
            run_summary_df = run_summary_df[
                run_summary_df["Database"].astype(str).str.upper() != "TOTAL"
            ].reset_index(drop=True)

            total_before = 0
            total_after = 0

            # for i in range(len(run_summary_df)):

                # db_name = str(run_summary_df.at[i, "Database"]).strip().upper()

                # if db_name not in after_counts:
                #     continue

                # try:
                #     before = int(run_summary_df.at[i, "Before"])
                # except:
                #     before = 0

                # after = after_counts[db_name]
                # excluded = before - after

                # run_summary_df.at[i, "After_Filter"] = after
                # run_summary_df.at[i, "Excluded"] = excluded

                # total_before += before
                # total_after += after


            for i in range(len(run_summary_df)):

                db_name = str(run_summary_df.at[i, "Database"]).strip().upper()

                if db_name not in after_counts:
                    continue

                try:
                    before = int(run_summary_df.at[i, "Before"])
                except:
                    before = 0

                after = after_counts[db_name]
                excluded = before - after

                run_summary_df.at[i, "After_Filter"] = after
                run_summary_df.at[i, "Excluded"] = excluded

                # Exclude the Variations sheet from TOTAL calculations
                if db_name != "VARIATIONS":
                    total_before += before
                    total_after += after

                    
            total_excluded = total_before - total_after

            run_summary_df.loc[len(run_summary_df)] = [
                "TOTAL",
                total_before,
                total_after,
                total_excluded
            ]

            print(f"📊 Total Before   : {total_before}")
            print(f"📊 Total After    : {total_after}")
            print(f"📊 Total Excluded : {total_excluded}")

            run_summary_df.to_excel(
                writer,
                sheet_name="Run_Summary",
                index=False
            )
    print(f"\n🎉 DONE: {output_file}")

    # ============================================
    # APPLY HIGHLIGHTING
    # ============================================
    if output_file.endswith("_Integrated_Somamut.xlsx"):

        print("🎨 Applying somatic highlighting (ALL sheets)...")
        highlight_excel_file(output_file, somatic_genes)

        print(f"🎨 Applying disease highlighting ONLY on {SHEET_NAME}...")
        apply_disease_highlighting(
            output_file,
            SHEET_NAME,
            disease_gene_sets   # ✅ FIXED
        )

        print("Applying the filters")
        run_all_filters(output_file)

        # ============================================
        # RESTORE ORIGINAL (PRE GERMLINE-FILTER) VARIATIONS SHEET
        # Runs last so nothing downstream touches it again.
        # ============================================
        print(f"🔄 Restoring original Variations sheet for sample {sample_folder}...")
        restore_original_variations(output_file, sample_folder)


print("\n==============================")
print("   SOMAMUT COMPLETED          ")
print("==============================\n")
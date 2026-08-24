import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill



# ======================================================
# LOAD SOMATIC GENES (GLOBAL)
# ======================================================
def load_somatic_genes(file_path):

    df = pd.read_excel(file_path, dtype=str)
    df.columns = df.columns.str.strip()

    col_name = "Unique (Somatic) - 1332 genes"

    if col_name not in df.columns:
        raise ValueError(f"❌ Column '{col_name}' not found")

    return set(
        df[col_name]
        .dropna()
        .astype(str)
        .str.strip()
        .str.upper()
    )


# ======================================================
# LOAD DISEASE-SPECIFIC GENE SETS
# ======================================================
def load_disease_gene_sets(file_path):

    df = pd.read_excel(file_path, dtype=str)
    df.columns = df.columns.str.strip()

    disease_cols = [
        "Cardiac - 565 genes",
        "Diabetes_pharmgkb_mito genes - 318 genes",
        "NDD - 890 genes",
        "Endometriosis - 53 genes",
        "Neurodevelopmental - 1733 genes",
    ]

    gene_sets = {}

    for col in disease_cols:
        if col in df.columns:
            gene_sets[col] = set(
                df[col]
                .dropna()
                .astype(str)
                .str.strip()
                .str.upper()
            )

    return gene_sets


# ======================================================
# EXISTING HIGHLIGHTING (ALL SHEETS)
# ======================================================
def highlight_excel_file(file_path, gene_set):

    wb = load_workbook(file_path)
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    for sheet_name in wb.sheetnames:

        ws = wb[sheet_name]

        if ws.max_row < 2:
            continue

        header = [cell.value for cell in ws[1]]

        symbol_cols = []
        if "SYMBOL" in header:
            symbol_cols.append(header.index("SYMBOL") + 1)
        if "SYMBOL_X" in header:
            symbol_cols.append(header.index("SYMBOL_X") + 1)

        if not symbol_cols:
            continue

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):

            for col_idx in symbol_cols:
                val = ws.cell(row=row[0].row, column=col_idx).value
                gene = str(val).strip().upper() if val else ""

                if gene in gene_set:
                    for cell in row:
                        cell.fill = red
                    break

    wb.save(file_path)


# ======================================================
# DISEASE HIGHLIGHTING (ONLY TARGET SHEET)
# ======================================================
def apply_disease_highlighting(file_path, target_sheet, disease_gene_sets):

    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    wb = load_workbook(file_path)

    if target_sheet not in wb.sheetnames:
        return

    ws = wb[target_sheet]

    header = [c.value for c in ws[1]]

    symbol_cols = []
    if "SYMBOL" in header:
        symbol_cols.append(header.index("SYMBOL") + 1)
    if "SYMBOL_X" in header:
        symbol_cols.append(header.index("SYMBOL_X") + 1)

    fills = {
        "Cardiac - 565 genes": PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid"),
        "Diabetes_pharmgkb_mito genes - 318 genes": PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid"),
        "NDD - 890 genes": PatternFill(start_color="C0C0C0", end_color="C0C0C0", fill_type="solid"),
        "Endometriosis - 53 genes": PatternFill(start_color="FF69B4", end_color="FF69B4", fill_type="solid"),
        "Neurodevelopmental - 1733 genes": PatternFill(start_color="87CEEB", end_color="87CEEB", fill_type="solid"),
    }

    multi_fill = PatternFill(start_color="E6CCFF", end_color="E6CCFF", fill_type="solid")

    match_col = len(header) + 1
    ws.cell(row=1, column=match_col, value="MATCH_SOURCE")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):

        matched = set()

        for col_idx in symbol_cols:
            gene = str(ws.cell(row=row[0].row, column=col_idx).value).strip().upper()

            for d, gset in disease_gene_sets.items():
                if gene in gset:
                    matched.add(d)

        if matched:
            ws.cell(row=row[0].row, column=match_col).value = ", ".join(sorted(matched))

        if len(matched) == 1:
            fill = fills[list(matched)[0]]
        elif len(matched) > 1:
            fill = multi_fill
        else:
            continue

        for cell in row:
            cell.fill = fill

    wb.save(file_path)


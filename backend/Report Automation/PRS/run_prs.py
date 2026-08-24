import os
from pathlib import Path

import pandas as pd # type: ignore
from openpyxl.styles import PatternFill # type: ignore


# ============================================================
# CONFIG
# ============================================================

INPUT_DIR = "Input"
OUTPUT_DIR = "output"

SOURCE_SHEET = "Sheet1"

PRS_COLUMN = "Polygenic Risk Score"

# Exact trait mappings
TRAIT_MAPPING = {
    "Cardiac": [
        "Heart Amyloid Deposition Measurement",
        "Congenital Left-Sided Heart Lesions",
    ],
    "Diabetes": [
        "Type 1 Diabetes Mellitus",
        "Type 2 Diabetes Mellitus",
    ],
    "Asthma": [
        "Asthma",
    ],
    "ND": [
        "Autism",
    ],
}

# NDD special rules
NDD_EXACT = [
    "Alzheimer Disease",
    "Parkinson Disease",
]

# Exact case-sensitive match
ND_CASE_SENSITIVE = [
    "Down Syndrome",
]

# Highlight highest PRS rows
HIGHLIGHT_FILL = PatternFill(
    fill_type="solid",
    start_color="C6EFCE",
    end_color="C6EFCE"
)


# ============================================================
# HELPERS
# ============================================================

def normalize_text(value):
    """Safe lowercase comparison."""
    if pd.isna(value):
        return ""
    return str(value).strip().lower()


def process_file(file_path):
    print(f"\nProcessing: {file_path.name}")

    # --------------------------------------------------------
    # Read file
    # --------------------------------------------------------
    df = pd.read_excel(file_path, sheet_name=SOURCE_SHEET)

    # Clean column names
    df.columns = df.columns.astype(str).str.strip()

    required_cols = ["Score Type", "Trait"]

    missing = [c for c in required_cols if c not in df.columns]

    if missing:
        raise ValueError(
            f"{file_path.name}: Missing required column(s): {missing}"
        )

    # --------------------------------------------------------
    # Filter Score Type = OR
    # --------------------------------------------------------
    score_type = (
        df["Score Type"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    filtered_df = df[score_type == "OR"].copy()

    print(f"Rows after OR filter: {len(filtered_df)}")

    # --------------------------------------------------------
    # Create buckets
    # --------------------------------------------------------
    sheet_data = {
        "Cardiac": [],
        "Diabetes": [],
        "NDD": [],
        "Asthma": [],
        "ND": [],
    }

    unmatched_count = 0

    for _, row in filtered_df.iterrows():

        trait_raw = row["Trait"]

        if pd.isna(trait_raw):
            unmatched_count += 1
            continue

        trait = str(trait_raw).strip()
        trait_lower = trait.lower()

        matched = False

        # ----------------------------------------------------
        # Cardiac / Diabetes / Asthma / Autism
        # ----------------------------------------------------
        for sheet_name, trait_list in TRAIT_MAPPING.items():

            normalized_traits = {
                x.strip().lower()
                for x in trait_list
            }

            if trait_lower in normalized_traits:
                sheet_data[sheet_name].append(row)
                matched = True
                break

        # ----------------------------------------------------
        # NDD
        # ----------------------------------------------------
        if not matched:

            ndd_exact = {
                x.lower()
                for x in NDD_EXACT
            }

            if (
                trait_lower in ndd_exact
                or "dementia" in trait_lower
            ):
                sheet_data["NDD"].append(row)
                matched = True

        # ----------------------------------------------------
        # ND (Down Syndrome exact case-sensitive)
        # ----------------------------------------------------
        if not matched:

            if trait in ND_CASE_SENSITIVE:
                sheet_data["ND"].append(row)
                matched = True

        if not matched:
            unmatched_count += 1

    # --------------------------------------------------------
    # Output file name
    # --------------------------------------------------------
    sample_name = file_path.stem

    if sample_name.endswith("_Merged"):
        sample_name = sample_name[:-7]

    output_file = (
        Path(OUTPUT_DIR)
        / f"{sample_name}_trait_processed.xlsx"
    )

    # --------------------------------------------------------
    # Write workbook
    # --------------------------------------------------------
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:

        # Summary FIRST
        summary_df = pd.DataFrame({
            "Sheet": [
                "Cardiac",
                "Diabetes",
                "NDD",
                "Asthma",
                "ND",
                "Total OR Rows Processed",
                "Unmatched Rows",
            ],
            "Count": [
                len(sheet_data["Cardiac"]),
                len(sheet_data["Diabetes"]),
                len(sheet_data["NDD"]),
                len(sheet_data["Asthma"]),
                len(sheet_data["ND"]),
                len(filtered_df),
                unmatched_count,
            ],
        })

        summary_df.to_excel(
            writer,
            sheet_name="Summary",
            index=False
        )

        # ----------------------------------------------------
        # Category sheets
        # ----------------------------------------------------
        for sheet_name in [
            "Cardiac",
            "Diabetes",
            "NDD",
            "Asthma",
            "ND",
        ]:

            rows = sheet_data[sheet_name]

            if rows:
                out_df = pd.DataFrame(rows)
            else:
                out_df = pd.DataFrame(columns=filtered_df.columns)

            out_df.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False
            )

            # ------------------------------------------------
            # Highlight highest PRS row(s)
            # ------------------------------------------------
            if (
                not out_df.empty
                and PRS_COLUMN in out_df.columns
            ):

                prs_series = pd.to_numeric(
                    out_df[PRS_COLUMN]
                    .astype(str)
                    .str.replace("%", "", regex=False)
                    .str.strip(),
                    errors="coerce"
                )

                if prs_series.notna().any():

                    max_score = prs_series.max()

                    ws = writer.sheets[sheet_name]

                    # Excel row 1 = header
                    for excel_row, score in enumerate(
                        prs_series,
                        start=2
                    ):

                        if (
                            pd.notna(score)
                            and score == max_score
                        ):

                            for cell in ws[excel_row]:
                                cell.fill = HIGHLIGHT_FILL

    print(f"Saved: {output_file.name}")
    print(f"Unmatched rows: {unmatched_count}")


# ============================================================
# MAIN
# ============================================================

def main():

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    input_files = sorted(
        Path(INPUT_DIR).glob("*.xlsx")
    )

    if not input_files:
        print("No .xlsx files found in input folder.")
        return

    for file_path in input_files:
        try:
            process_file(file_path)
        except Exception as e:
            print(f"ERROR processing {file_path.name}: {e}")


if __name__ == "__main__":
    main()
